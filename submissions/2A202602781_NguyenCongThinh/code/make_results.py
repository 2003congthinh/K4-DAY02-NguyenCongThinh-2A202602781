"""make_results.py - Bước 5: tổng hợp mọi kết quả thành results.xlsx, bảng markdown và biểu đồ tổng hợp.

Chạy từ thư mục bài nộp:  python code/make_results.py      (thường gọi qua run_all.py results)
Đọc: runs/*/seed*/summary.json, results/*.csv|json, predictions/*.csv
Ghi: results.xlsx (sheet Backbones, Training, Inference, Final, PerClass, Latency, Summary - GUIDE 6.1),
     results/tables.md (các bảng dạng markdown để dán vào report.md),
     figures/backbone_tradeoff.png, figures/training_deltas.png
Mọi chỉ số test được tính lại từ predictions/ bằng hàm của eval.py gốc, nên khớp với `eval.py score`.
"""
from __future__ import annotations

import glob
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))

import experiments as ex  # noqa: E402
import train  # noqa: E402

sys.path.insert(0, str(train.REPO_ROOT))
import eval as ev  # noqa: E402

CHINEE, SNAKE = 0, 7


def _summary(key: str) -> dict | None:
    """summary.json của run theo khoá, hoặc None nếu chưa chạy."""
    try:
        f = train.run_dir(ex.get_cfg(key)) / "summary.json"
    except RuntimeError:
        return None
    return json.loads(f.read_text(encoding="utf-8")) if f.exists() else None


def sheet_backbones() -> pd.DataFrame:
    """Sheet Backbones (GUIDE 6.1). Output: DataFrame một dòng mỗi B0x."""
    lat = pd.read_csv("results/backbone_latency.csv").set_index("exp_id") if Path("results/backbone_latency.csv").exists() else None
    chosen = ex.load_decisions().get("chosen_backbone")
    rows = []
    for k in ex.BACKBONES:
        s = _summary(k)
        if s is None:
            continue
        rows.append({"exp_id": k, "backbone": s["backbone"], "weight tag (timm)": s["weight_tag"],
                     "params (M)": s["params_m"], "GMAC": s["gmacs"], "img size (px)": s["img_size"],
                     "epochs": s["epochs"], "seed": s["seed"], "macro-F1 val": s["val_macro_f1"],
                     "top-1 val": s["val_top1"], "best epoch": s["best_epoch"],
                     "train time / epoch (s)": s["train_time_per_epoch_s"],
                     "latency batch-1 p50 (ms)": lat.loc[k, "p50_ms"] if lat is not None and k in lat.index else np.nan,
                     "device": s["device"],
                     "note": "CHỌN đi tiếp" if k == chosen else ""})
    return pd.DataFrame(rows)


def sheet_training() -> pd.DataFrame:
    """Sheet Training: T00..T12, Δ so với T00, F1 lớp hiếm (val, seed 0)."""
    dec = ex.load_decisions()
    base = _summary("T00")
    if base is None:
        return pd.DataFrame()
    rows = []
    table = {"T00": ("-", "công thức nền", {}), **ex.TRAINING,
             "T12": ("combo", "kết hợp: " + ", ".join(dec.get("combo_from", [])), dec.get("combo", {}))}
    for k, (axis, desc, ov) in table.items():
        s = _summary(k)
        if s is None:
            continue
        rows.append({"exp_id": k, "backbone": s["backbone"], "axis (A-F)": axis, "khác T00 ở": desc,
                     "overrides": json.dumps(ov, ensure_ascii=False), "seed": s["seed"],
                     "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"],
                     "Δ macro-F1 vs T00": s["val_macro_f1"] - base["val_macro_f1"],
                     "F1 Chinee apple val": s["val_f1_per_class"][CHINEE], "F1 Snake weed val": s["val_f1_per_class"][SNAKE],
                     "best epoch": s["best_epoch"], "train time / epoch (s)": s["train_time_per_epoch_s"],
                     "note": ("cùng run với " + s["reused_from"]) if s.get("reused_from") else "1 seed"})
    return pd.DataFrame(rows)


def sheet_inference() -> pd.DataFrame:
    """Sheet Inference từ results/inference.csv (đổi tên cột kèm đơn vị)."""
    f = Path("results/inference.csv")
    if not f.exists():
        return pd.DataFrame()
    df = pd.read_csv(f)
    return df.rename(columns={"method": "phương pháp", "model": "mô hình/checkpoint", "macro_f1_val": "macro-F1 val",
                              "top1_val": "top-1 val", "ece_val": "ECE val", "f1_chinee_val": "F1 Chinee apple val",
                              "f1_snake_val": "F1 Snake weed val", "p50_ms": "p50 batch-1 (ms)",
                              "p95_ms": "p95 batch-1 (ms)", "p99_ms": "p99 batch-1 (ms)",
                              "img_per_s_b32": "throughput batch-32 (ảnh/s)", "rel_cost_vs_I00": "chi phí tương đối vs I00"})


def _group(pattern: str, ref_csv: str | None):
    """eval.load_group nếu có file khớp, ngược lại None."""
    return ev.load_group(pattern, ref_csv) if glob.glob(pattern) else None


def sheet_final() -> tuple[pd.DataFrame, dict]:
    """Sheet Final: từng seed + dòng mean ± std cho F01, mốc T00, F01 chưa TS.

    Output: (DataFrame, groups) với groups = {"F01": eval.Group, "T00": ..., "F01_uncal": ...} để dùng lại.
    """
    lab = Path(ex.BASE.labels_dir)
    dec = ex.load_decisions()
    inf = dec.get("inference", {})
    desc = {"F01": f"{ex.chosen_backbone()[0]} + {dec.get('final_source')} recipe "
                   f"{json.dumps(dec.get('final_recipe', {}), ensure_ascii=False)} + {inf.get('exp_id')} "
                   f"({inf.get('method')}, {inf.get('space')}) + temperature scaling",
            "T00": f"{ex.chosen_backbone()[0]} + công thức nền T00 + 1 view (I00), không TS",
            "F01_uncal": "F01 trước temperature scaling"}
    rows, groups = [], {}
    for tag in ("F01", "T00", "F01_uncal"):
        g = _group(f"predictions/{tag}_seed*_test.csv", str(lab / "test_subset0.csv"))
        if g is None:
            continue
        groups[tag] = g
        gv = _group(f"predictions/{tag}_seed*_val.csv", str(lab / "val_subset0.csv"))
        val_f1 = {p.seed: m["macro_f1"] for p, m in zip(gv.preds, gv.metrics)} if gv else {}
        for p, m in zip(g.preds, g.metrics):
            rows.append({"exp_id": tag, "cấu hình": desc[tag], "seed": p.seed, "macro-F1 val": val_f1.get(p.seed, np.nan),
                         "macro-F1 test": m["macro_f1"], "top-1 test": m["top1"], "balanced acc test": m["balanced_acc"],
                         "ECE test": m["ece"], "recall Chinee apple test": m["recall"][CHINEE],
                         "recall Snake weed test": m["recall"][SNAKE]})
        vm = ev.fmt(*ev.mean_std(list(val_f1.values()))) if val_f1 else ""
        summ = {k: g.summary[k] for k in ("macro_f1", "top1", "balanced_acc", "ece")}
        rows.append({"exp_id": f"{tag} mean ± std", "cấu hình": f"{len(g.preds)} seed, std mẫu ddof=1", "seed": "all",
                     "macro-F1 val": vm, "macro-F1 test": ev.fmt(*summ["macro_f1"]),
                     "top-1 test": ev.fmt(*summ["top1"]), "balanced acc test": ev.fmt(*summ["balanced_acc"]),
                     "ECE test": ev.fmt(*summ["ece"]),
                     "recall Chinee apple test": ev.fmt(g.summary["recall"][0][CHINEE], g.summary["recall"][1][CHINEE]),
                     "recall Snake weed test": ev.fmt(g.summary["recall"][0][SNAKE], g.summary["recall"][1][SNAKE])})
    return pd.DataFrame(rows), groups


def sheet_perclass(groups: dict) -> pd.DataFrame:
    """Sheet PerClass: precision/recall/F1 (mean ± std qua seed) từng lớp, test."""
    names = ev.load_names(str(Path(ex.BASE.labels_dir) / "labels.csv"))
    rows = []
    for tag in ("F01", "T00"):
        g = groups.get(tag)
        if g is None:
            continue
        for i, n in enumerate(names):
            rows.append({"cấu hình": tag, "lớp": n, "số ảnh test": int(g.metrics[0]["support"][i]),
                         "precision": ev.fmt(g.summary["precision"][0][i], g.summary["precision"][1][i]),
                         "recall": ev.fmt(g.summary["recall"][0][i], g.summary["recall"][1][i]),
                         "F1": ev.fmt(g.summary["f1"][0][i], g.summary["f1"][1][i])})
    return pd.DataFrame(rows)


def sheet_latency() -> pd.DataFrame:
    """Sheet Latency: mọi điều kiện đo (Bước 1 sơ bộ, Bước 3, pipeline chung kết)."""
    parts = []
    if Path("results/latency.csv").exists():
        parts.append(pd.read_csv("results/latency.csv"))
    if Path("results/backbone_latency.csv").exists():
        b = pd.read_csv("results/backbone_latency.csv")
        parts.append(pd.DataFrame({"config": b["exp_id"] + " " + b["backbone"] + " (Bước 1, sơ bộ)", "gpu": b["device"],
                                   "dtype": "fp32", "batch": 1, "bn_fused": False, "p50": b["p50_ms"],
                                   "p95": b["p95_ms"], "p99": np.nan, "img_per_s": 1000 / b["p50_ms"]}))
    if Path("results/final_temperatures.json").exists():
        l = json.loads(Path("results/final_temperatures.json").read_text(encoding="utf-8"))["latency_batch1"]
        parts.append(pd.DataFrame([{"config": l["pipeline"] + " (chung kết)", "gpu": l["device"], "dtype": "fp32",
                                    "batch": 1, "bn_fused": False, "p50": l["p50"], "p95": l["p95"], "p99": l["p99"],
                                    "img_per_s": 1000 / l["p50"]}]))
    if not parts:
        return pd.DataFrame()
    return pd.concat(parts, ignore_index=True).rename(columns={"config": "cấu hình", "gpu": "GPU / thiết bị",
                                                              "bn_fused": "gộp BN", "p50": "p50 (ms)", "p95": "p95 (ms)",
                                                              "p99": "p99 (ms)", "img_per_s": "ảnh/s"})


def sheet_summary(bb: pd.DataFrame, tr: pd.DataFrame, infr: pd.DataFrame, fin: pd.DataFrame, groups: dict) -> pd.DataFrame:
    """Sheet Summary: top 10 theo macro-F1 val + so sánh chung kết vs mốc trên test."""
    rows = []
    for _, r in bb.iterrows():
        rows.append({"exp_id": r["exp_id"], "loại": "backbone", "mô tả": r["backbone"], "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "p50 batch-1 (ms)": r["latency batch-1 p50 (ms)"]})
    for _, r in tr.iterrows():
        rows.append({"exp_id": r["exp_id"], "loại": "training", "mô tả": r["khác T00 ở"], "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "p50 batch-1 (ms)": np.nan})
    for _, r in infr.dropna(subset=["macro-F1 val"]).iterrows():
        rows.append({"exp_id": r["exp_id"], "loại": "inference", "mô tả": r["phương pháp"], "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "p50 batch-1 (ms)": r["p50 batch-1 (ms)"]})
    top = pd.DataFrame(rows).sort_values("macro-F1 val", ascending=False).head(10).reset_index(drop=True)
    top.insert(0, "hạng", range(1, len(top) + 1))
    if "F01" in groups and "T00" in groups:
        f, b = groups["F01"].summary, groups["T00"].summary
        delta = f["macro_f1"][0] - b["macro_f1"][0]
        s = max(v for v in (f["macro_f1"][1], b["macro_f1"][1]) if np.isfinite(v))
        extra = pd.DataFrame([
            {"hạng": "", "exp_id": "TEST", "loại": "chung kết vs mốc", "mô tả": "F01 (chung kết)",
             "macro-F1 val": ev.fmt(*f["macro_f1"]) + " (test)", "top-1 val": ev.fmt(*f["top1"]) + " (test)"},
            {"hạng": "", "exp_id": "TEST", "loại": "chung kết vs mốc", "mô tả": "T00 + I00 (mốc)",
             "macro-F1 val": ev.fmt(*b["macro_f1"]) + " (test)", "top-1 val": ev.fmt(*b["top1"]) + " (test)"},
            {"hạng": "", "exp_id": "TEST", "loại": "Δ", "mô tả": f"Δ macro-F1 = {delta:+.4f}, std lớn hơn s = {s:.4f}",
             "macro-F1 val": "vượt nhiễu" if delta > s else "KHÔNG vượt nhiễu"}])
        top = pd.concat([top, extra], ignore_index=True)
    return top


def write_xlsx(sheets: dict[str, pd.DataFrame], path: Path) -> None:
    """Ghi xlsx có định dạng: freeze tiêu đề, 4 chữ số, độ rộng cột, tô dòng tốt nhất.

    Input : sheets (dict tên sheet -> DataFrame), path. Output: None.
    """
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    hi = PatternFill("solid", fgColor="FFF2CC")
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = "B2"
            for c in ws[1]:
                c.font = Font(bold=True)
            for j, col in enumerate(df.columns, start=1):
                width = max([len(str(col))] + [len(f"{v:.4f}" if isinstance(v, float) else str(v)) for v in df[col]][:200])
                ws.column_dimensions[get_column_letter(j)].width = min(60, width + 2)
                for i in range(2, len(df) + 2):
                    cell = ws.cell(i, j)
                    if isinstance(cell.value, float):
                        cell.number_format = "0.0000"
            col = "macro-F1 val" if "macro-F1 val" in df.columns else None
            if col and len(df):
                numeric = pd.to_numeric(df[col], errors="coerce")
                if numeric.notna().any():
                    best = int(numeric.idxmax()) + 2
                    for c in ws[best]:
                        c.fill = hi
                        c.font = Font(bold=True)


def figures(bb: pd.DataFrame, tr: pd.DataFrame) -> None:
    """Biểu đồ tổng hợp: backbone (F1 theo độ trễ, cỡ chấm = params) và Δ của T."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    Path("figures").mkdir(exist_ok=True)
    if len(bb):
        fig, ax = plt.subplots(figsize=(7, 5))
        ax.scatter(bb["latency batch-1 p50 (ms)"], bb["macro-F1 val"], s=bb["params (M)"] * 12, alpha=0.6)
        for _, r in bb.iterrows():
            ax.annotate(f"{r['exp_id']} {r['backbone']}\n{r['params (M)']:.1f}M, {r['GMAC']:.2f} GMAC",
                        (r["latency batch-1 p50 (ms)"], r["macro-F1 val"]), fontsize=7, xytext=(5, 3), textcoords="offset points")
        ax.set_xlabel("độ trễ p50 batch 1 (ms)"); ax.set_ylabel("macro-F1 val (seed 0)")
        ax.set_title("Backbone: chất lượng theo độ trễ (cỡ chấm ~ số tham số)"); ax.grid(alpha=0.3)
        fig.tight_layout(); fig.savefig("figures/backbone_tradeoff.png", dpi=110); plt.close(fig)
    if len(tr):
        d = tr[tr["exp_id"] != "T00"]
        fig, ax = plt.subplots(figsize=(8, 4.5))
        colors = ["#55A868" if v > 0 else "#C44E52" for v in d["Δ macro-F1 vs T00"]]
        ax.barh(d["exp_id"] + " " + d["khác T00 ở"].str.slice(0, 38), d["Δ macro-F1 vs T00"], color=colors)
        ax.axvline(0, color="k", lw=0.8)
        ax.set_xlabel("Δ macro-F1 val so với T00 (1 seed)"); ax.set_title("Ablation công thức huấn luyện")
        ax.invert_yaxis(); ax.grid(alpha=0.3, axis="x")
        fig.tight_layout(); fig.savefig("figures/training_deltas.png", dpi=110); plt.close(fig)


def to_md(df: pd.DataFrame) -> str:
    """DataFrame -> bảng markdown (số thực 4 chữ số). Output: str."""
    if df.empty:
        return "_(chưa có dữ liệu)_"
    d = df.copy()
    for c in d.columns:
        d[c] = [f"{v:.4f}" if isinstance(v, float) and np.isfinite(v) else ("" if isinstance(v, float) else str(v)) for v in d[c]]
    lines = ["| " + " | ".join(map(str, d.columns)) + " |", "|" + "---|" * len(d.columns)]
    lines += ["| " + " | ".join(r) + " |" for r in d.astype(str).values.tolist()]
    return "\n".join(lines)


def main() -> None:
    """Tạo results.xlsx + results/tables.md + figures. Input/Output: không/None."""
    bb, tr, infr = sheet_backbones(), sheet_training(), sheet_inference()
    fin, groups = sheet_final()
    pc, lat = sheet_perclass(groups), sheet_latency()
    summ = sheet_summary(bb, tr, infr, fin, groups)
    sheets = {"Backbones": bb, "Training": tr, "Inference": infr, "Final": fin, "PerClass": pc, "Latency": lat,
              "Summary": summ}
    write_xlsx(sheets, Path("results.xlsx"))
    figures(bb, tr)
    md = [f"# Bảng kết quả (tự sinh bởi code/make_results.py, profile {ex.PROFILE})\n"]
    md += [f"## {k}\n\n{to_md(v)}\n" for k, v in sheets.items()]
    md.append("## Lựa chọn giữa các bước (results/decisions.json)\n\n```json\n"
              + json.dumps(ex.load_decisions(), indent=2, ensure_ascii=False) + "\n```\n")
    Path("results/tables.md").write_text("\n".join(md), encoding="utf-8")
    print("[make_results] đã ghi results.xlsx, results/tables.md, figures/backbone_tradeoff.png, figures/training_deltas.png")


if __name__ == "__main__":
    main()
